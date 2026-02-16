import os
from datetime import datetime
from typing import List, Dict
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from django.conf import settings
from django.utils import timezone


class ExcelReportGenerator:
    """Generate comprehensive Excel reports for security assessments"""
    
    # Color scheme
    HEADER_COLOR = "366092"
    CRITICAL_COLOR = "C00000"
    HIGH_COLOR = "FF6B6B"
    MEDIUM_COLOR = "FFA500"
    LOW_COLOR = "FFD700"
    INFO_COLOR = "90EE90"
    
    def __init__(self):
        self.wb = Workbook()
        self.border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
    
    @staticmethod
    def _make_naive(dt):
        """Convert timezone-aware datetime to naive datetime for Excel compatibility"""
        if dt is None:
            return None
        if isinstance(dt, datetime):
            if timezone.is_aware(dt):
                # Convert to local time and remove timezone info
                return timezone.localtime(dt).replace(tzinfo=None)
        return dt
    
    def generate_report(
        self,
        assessment: Dict,
        vulnerabilities: List[Dict],
        history: List[Dict] = None
    ) -> str:
        """
        Generate comprehensive Excel report
        
        Returns:
            File path to generated report
        """
        
        # Remove default sheet
        if 'Sheet' in self.wb.sheetnames:
            del self.wb['Sheet']
        
        # Create sheets
        self._create_summary_sheet(assessment, vulnerabilities)
        self._create_findings_sheet(vulnerabilities)
        self._create_remediation_sheet(vulnerabilities)
        if history:
            self._create_history_sheet(history)
        
        # Save file
        current_time = self._make_naive(timezone.now())
        filename = f"security_assessment_{assessment['id']}_{current_time.strftime('%Y%m%d_%H%M%S')}.xlsx"
        filepath = os.path.join(settings.MEDIA_ROOT, 'reports', filename)
        os.makedirs(os.path.dirname(filepath), exist_ok=True)
        
        self.wb.save(filepath)
        return filepath
    
    def _create_summary_sheet(self, assessment: Dict, vulnerabilities: List[Dict]):
        """Create executive summary sheet"""
        
        ws = self.wb.create_sheet("Executive Summary", 0)
        
        # Title
        ws['A1'] = "Security Assessment Report"
        ws['A1'].font = Font(size=18, bold=True, color="FFFFFF")
        ws['A1'].fill = PatternFill(start_color=self.HEADER_COLOR, fill_type="solid")
        ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
        ws.merge_cells('A1:F1')
        ws.row_dimensions[1].height = 30
        
        # Assessment info
        row = 3
        current_time = self._make_naive(timezone.now())
        info_fields = [
            ("Assessment ID:", assessment.get('id', 'N/A')),
            ("Date:", current_time.strftime('%Y-%m-%d %H:%M')),
            ("Status:", assessment.get('status', 'N/A').upper()),
            ("Overall Risk Score:", f"{assessment.get('overall_risk_score', 0)}/100"),
        ]
        
        for label, value in info_fields:
            ws[f'A{row}'] = label
            ws[f'A{row}'].font = Font(bold=True)
            ws[f'B{row}'] = value
            ws.merge_cells(f'B{row}:F{row}')
            row += 1
        
        # Risk score visualization
        row += 1
        ws[f'A{row}'] = "Risk Level:"
        ws[f'A{row}'].font = Font(bold=True)
        
        risk_score = assessment.get('overall_risk_score', 0)
        risk_level, risk_color = self._get_risk_level(risk_score)
        ws[f'B{row}'] = risk_level
        ws[f'B{row}'].fill = PatternFill(start_color=risk_color, fill_type="solid")
        ws[f'B{row}'].font = Font(bold=True, color="FFFFFF")
        ws[f'B{row}'].alignment = Alignment(horizontal='center')
        
        # Vulnerability summary
        row += 2
        ws[f'A{row}'] = "Vulnerability Summary"
        ws[f'A{row}'].font = Font(size=14, bold=True, color="FFFFFF")
        ws[f'A{row}'].fill = PatternFill(start_color=self.HEADER_COLOR, fill_type="solid")
        ws.merge_cells(f'A{row}:F{row}')
        
        row += 1
        severity_counts = self._count_by_severity(vulnerabilities)
        
        headers = ['Severity', 'Count', 'Open', 'Fixed', 'In Progress', 'Accepted']
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=row, column=col)
            cell.value = header
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color="D9D9D9", fill_type="solid")
            cell.alignment = Alignment(horizontal='center')
            cell.border = self.border
        
        row += 1
        for severity in ['Critical', 'High', 'Medium', 'Low', 'Informational']:
            counts = severity_counts.get(severity.lower(), {})
            color = getattr(self, f"{severity.upper()}_COLOR", "FFFFFF")
            
            ws.cell(row=row, column=1, value=severity)
            ws.cell(row=row, column=1).fill = PatternFill(start_color=color, fill_type="solid")
            ws.cell(row=row, column=1).font = Font(bold=True)
            
            ws.cell(row=row, column=2, value=counts.get('total', 0))
            ws.cell(row=row, column=3, value=counts.get('open', 0))
            ws.cell(row=row, column=4, value=counts.get('fixed', 0))
            ws.cell(row=row, column=5, value=counts.get('in_progress', 0))
            ws.cell(row=row, column=6, value=counts.get('accepted', 0))
            
            for col in range(1, 7):
                ws.cell(row=row, column=col).border = self.border
                ws.cell(row=row, column=col).alignment = Alignment(horizontal='center')
            
            row += 1
        
        # Risk reasoning
        row += 2
        ws[f'A{row}'] = "Risk Assessment Reasoning"
        ws[f'A{row}'].font = Font(size=14, bold=True, color="FFFFFF")
        ws[f'A{row}'].fill = PatternFill(start_color=self.HEADER_COLOR, fill_type="solid")
        ws.merge_cells(f'A{row}:F{row}')
        
        row += 1
        ws[f'A{row}'] = assessment.get('risk_reasoning', 'N/A')
        ws[f'A{row}'].alignment = Alignment(wrap_text=True, vertical='top')
        ws.merge_cells(f'A{row}:F{row + 5}')
        
        # Adjust column widths
        ws.column_dimensions['A'].width = 25
        for col in ['B', 'C', 'D', 'E', 'F']:
            ws.column_dimensions[col].width = 15
    
    def _create_findings_sheet(self, vulnerabilities: List[Dict]):
        """Create detailed findings sheet"""
        
        ws = self.wb.create_sheet("Findings")
        
        # Headers
        headers = [
            'ID', 'Control Title', 'Control Description', 'Control Impact',
            'Control Recommendation', 'Severity', 'Status', 'Affected Devices',
            'Category Tag', 'Framework Mapping', 'CVSS Score', 'CWE ID', 'OWASP Category'
        ]
        
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col)
            cell.value = header
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color=self.HEADER_COLOR, fill_type="solid")
            cell.alignment = Alignment(horizontal='center', wrap_text=True)
            cell.border = self.border
        
        # Data rows
        for row, vuln in enumerate(vulnerabilities, 2):
            # Get severity color
            severity = vuln.get('severity', 'low').lower()
            severity_color = getattr(self, f"{severity.upper()}_COLOR", "FFFFFF")
            
            data = [
                row - 1,
                vuln.get('control_title', ''),
                vuln.get('control_description', ''),
                vuln.get('control_impact', ''),
                vuln.get('control_recommendation', ''),
                vuln.get('severity', '').upper(),
                vuln.get('status', 'open').upper(),
                vuln.get('affected_devices', ''),
                vuln.get('category_tag', ''),
                vuln.get('framework_mapping', ''),
                vuln.get('cvss_score', ''),
                vuln.get('cwe_id', ''),
                vuln.get('owasp_category', '')
            ]
            
            for col, value in enumerate(data, 1):
                cell = ws.cell(row=row, column=col)
                cell.value = value
                cell.border = self.border
                cell.alignment = Alignment(wrap_text=True, vertical='top')
                
                # Color severity column
                if col == 6:  # Severity column
                    cell.fill = PatternFill(start_color=severity_color, fill_type="solid")
                    if severity in ['critical', 'high']:
                        cell.font = Font(bold=True, color="FFFFFF")
        
        # Adjust column widths
        column_widths = {
            'A': 8, 'B': 30, 'C': 50, 'D': 40, 'E': 40,
            'F': 12, 'G': 12, 'H': 25, 'I': 20, 'J': 30,
            'K': 12, 'L': 12, 'M': 25
        }
        for col, width in column_widths.items():
            ws.column_dimensions[col].width = width
        
        # Freeze header row
        ws.freeze_panes = 'A2'
    
    def _create_remediation_sheet(self, vulnerabilities: List[Dict]):
        """Create remediation tracking sheet"""
        
        ws = self.wb.create_sheet("Remediation Plan")
        
        # Headers
        headers = [
            'Finding ID', 'Control Title', 'Severity', 'Status',
            'Remediation Control', 'Implementation Details',
            'Risk Reduction %', 'Control Status', 'Verified By', 'Verified Date'
        ]
        
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col)
            cell.value = header
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color=self.HEADER_COLOR, fill_type="solid")
            cell.alignment = Alignment(horizontal='center', wrap_text=True)
            cell.border = self.border
        
        # Data rows
        row = 2
        for idx, vuln in enumerate(vulnerabilities, 1):
            remediation_controls = vuln.get('remediation_controls', [])
            
            if not remediation_controls:
                # Add empty row for tracking
                ws.cell(row=row, column=1, value=idx)
                ws.cell(row=row, column=2, value=vuln.get('control_title', ''))
                ws.cell(row=row, column=3, value=vuln.get('severity', '').upper())
                ws.cell(row=row, column=4, value=vuln.get('status', 'open').upper())
                
                for col in range(1, 11):
                    ws.cell(row=row, column=col).border = self.border
                row += 1
            else:
                for control in remediation_controls:
                    # Convert verified_at to naive datetime
                    verified_at = control.get('verified_at', '')
                    if verified_at:
                        verified_at = self._make_naive(verified_at)
                    
                    data = [
                        idx,
                        vuln.get('control_title', ''),
                        vuln.get('severity', '').upper(),
                        vuln.get('status', 'open').upper(),
                        control.get('control_name', ''),
                        control.get('implementation_details', ''),
                        control.get('risk_reduction_percentage', ''),
                        control.get('status', '').upper(),
                        control.get('verified_by', ''),
                        verified_at
                    ]
                    
                    for col, value in enumerate(data, 1):
                        cell = ws.cell(row=row, column=col)
                        cell.value = value
                        cell.border = self.border
                        cell.alignment = Alignment(wrap_text=True, vertical='top')
                    
                    row += 1
        
        # Adjust column widths
        column_widths = {
            'A': 10, 'B': 30, 'C': 12, 'D': 12, 'E': 30,
            'F': 40, 'G': 15, 'H': 15, 'I': 20, 'J': 15
        }
        for col, width in column_widths.items():
            ws.column_dimensions[col].width = width
        
        ws.freeze_panes = 'A2'
    
    def _create_history_sheet(self, history: List[Dict]):
        """Create risk score history sheet"""
        
        ws = self.wb.create_sheet("Score History")
        
        # Headers
        headers = ['Date', 'Previous Score', 'New Score', 'Change', 'Reason', 'Changed By']
        
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col)
            cell.value = header
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color=self.HEADER_COLOR, fill_type="solid")
            cell.alignment = Alignment(horizontal='center')
            cell.border = self.border
        
        # Data rows
        for row, entry in enumerate(history, 2):
            prev_score = entry.get('previous_score')
            new_score = entry.get('new_score')
            change = new_score - prev_score if prev_score is not None else 0
            
            # Convert timestamp to naive datetime for Excel
            timestamp = entry.get('timestamp', '')
            if timestamp:
                timestamp = self._make_naive(timestamp)
            
            data = [
                timestamp,
                prev_score if prev_score is not None else 'N/A',
                new_score,
                change,
                entry.get('change_reason', ''),
                entry.get('changed_by', '')
            ]
            
            for col, value in enumerate(data, 1):
                cell = ws.cell(row=row, column=col)
                cell.value = value
                cell.border = self.border
                cell.alignment = Alignment(wrap_text=True, vertical='top')
                
                # Color change column
                if col == 4 and isinstance(change, (int, float)):
                    if change < 0:
                        cell.fill = PatternFill(start_color="90EE90", fill_type="solid")
                    elif change > 0:
                        cell.fill = PatternFill(start_color="FFB6C1", fill_type="solid")
        
        # Adjust column widths
        ws.column_dimensions['A'].width = 20
        ws.column_dimensions['B'].width = 15
        ws.column_dimensions['C'].width = 15
        ws.column_dimensions['D'].width = 10
        ws.column_dimensions['E'].width = 50
        ws.column_dimensions['F'].width = 20
        
        ws.freeze_panes = 'A2'
    
    @staticmethod
    def _get_risk_level(score: int) -> tuple:
        """Get risk level and color based on score"""
        if score >= 81:
            return "HIGH RISK", "C00000"
        elif score >= 61:
            return "MEDIUM-HIGH RISK", "FF6B6B"
        elif score >= 41:
            return "MEDIUM RISK", "FFA500"
        elif score >= 21:
            return "LOW-MEDIUM RISK", "FFD700"
        else:
            return "LOW RISK", "90EE90"
    
    @staticmethod
    def _count_by_severity(vulnerabilities: List[Dict]) -> Dict:
        """Count vulnerabilities by severity and status"""
        counts = {}
        
        for vuln in vulnerabilities:
            severity = vuln.get('severity', 'low').lower()
            status = vuln.get('status', 'open').lower()
            
            if severity not in counts:
                counts[severity] = {'total': 0, 'open': 0, 'fixed': 0, 'in_progress': 0, 'accepted': 0}
            
            counts[severity]['total'] += 1
            counts[severity][status] = counts[severity].get(status, 0) + 1
        
        return counts